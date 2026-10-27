"""Build the final presentation (.pptx, 16:9) from results/.

    python reporting/build_slides.py --results results --out ICT4442_Final_Presentation.pptx

Python-pptx is used (not a JS tool) so the deck can be regenerated on Kaggle right after the runs.
Missing numbers show as highlighted TBD; nothing is invented.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
import content as C  # noqa: E402
from results import MODEL_ORDER, TBD, Results, num, pct  # noqa: E402

# Palette: asphalt + road-marking yellow
ASPHALT = RGBColor(0x23, 0x26, 0x2B)
ASPHALT2 = RGBColor(0x34, 0x38, 0x3F)
YELLOW = RGBColor(0xF2, 0xB7, 0x05)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x5F, 0x66, 0x70)
LIGHT = RGBColor(0xF1, 0xF3, 0xF5)
LINE = RGBColor(0xD5, 0xD9, 0xDE)
MODEL_COLORS = {"resnet50": RGBColor(0x2A, 0x6F, 0x97), "yolov8": RGBColor(0xE7, 0x6F, 0x51),
                "unet": RGBColor(0x2A, 0x9D, 0x8F), "vit_b16": RGBColor(0x8E, 0x5E, 0xA2)}
HEAD, BODY = "Cambria", "Calibri"
W, H = Inches(13.333), Inches(7.5)
PRESENTER = {"resnet50": "Lakshya", "yolov8": "Dhruv", "unet": "Kumar", "vit_b16": "Manas"}


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    # ---------------- primitives ----------------
    def slide(self, dark=False, notes=""):
        s = self.prs.slides.add_slide(self.blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = ASPHALT if dark else WHITE
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        self.n += 1
        if self.n > 1:
            self.text(s, f"{self.n}", Inches(12.4), Inches(7.0), Inches(0.6), Inches(0.3), 10,
                      color=MUTED if not dark else LINE, align=PP_ALIGN.RIGHT)
            self.text(s, C.SHORT_TITLE + "  ·  ICT-4442", Inches(0.6), Inches(7.0), Inches(6), Inches(0.3), 10,
                      color=MUTED if not dark else LINE)
        return s

    def text(self, s, text, x, y, w, h, size=16, bold=False, color=INK, font=BODY, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, italic=False):
        box = s.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, attr, 0)
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            self.runs(p, line, size, bold, color, font, italic)
        return box

    def runs(self, p, line, size, bold, color, font, italic=False):
        parts = str(line).split(TBD)
        for k, part in enumerate(parts):
            if part:
                r = p.add_run()
                r.text = part
                self.style(r, size, bold, color, font, italic)
            if k < len(parts) - 1:
                r = p.add_run()
                r.text = TBD
                self.style(r, size, True, RGBColor(0xB4, 0x23, 0x18), font)

    @staticmethod
    def style(r, size, bold, color, font, italic=False):
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color

    def bullets(self, s, items, x, y, w, h, size=16, color=INK, gap=8):
        box = s.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, attr, 0)
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(gap)
            bold_part, _, rest = item.partition("|") if "|" in item else ("", "", item)
            r = p.add_run()
            r.text = "▪  "
            self.style(r, size, True, YELLOW if color == WHITE else MODEL_COLORS["resnet50"], BODY)
            if bold_part:
                r = p.add_run()
                r.text = bold_part + " "
                self.style(r, size, True, color, BODY)
            self.runs(p, rest, size, False, color, BODY)
        return box

    def title(self, s, text, kicker=None, dark=False):
        if kicker:
            self.text(s, kicker.upper(), Inches(0.6), Inches(0.38), Inches(10), Inches(0.3), 12, True,
                      YELLOW if dark else MUTED)
        self.text(s, text, Inches(0.6), Inches(0.68), Inches(12.1), Inches(0.8), 32, True, WHITE if dark else INK, HEAD)

    def card(self, s, x, y, w, h, fill=LIGHT):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.adjustments[0] = 0.06
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
        shape.shadow.inherit = False
        return shape

    def badge(self, s, label, x, y, color, size=0.62):
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(size), Inches(size))
        c.fill.solid()
        c.fill.fore_color.rgb = color
        c.line.fill.background()
        tf = c.text_frame
        tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        self.style(r, 18, True, WHITE, HEAD)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    def picture(self, s, path, x, y, w=None, h=None, caption=None):
        if path is None or not Path(path).exists():
            box = self.card(s, x, y, w or Inches(5), h or Inches(3), LIGHT)
            box.line.color.rgb = LINE
            box.line.dash_style = 4
            self.text(s, "Figure appears here after the final notebook run", x, y, w or Inches(5), h or Inches(3), 13,
                      color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, italic=True)
            return
        from PIL import Image

        iw, ih = Image.open(path).size
        if w and h:  # fit inside box keeping aspect
            scale = min(w / iw, h / ih)
            pw, ph = int(iw * scale), int(ih * scale)
            pic = s.shapes.add_picture(str(path), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)
        else:
            pic = s.shapes.add_picture(str(path), x, y, w, h)
        if caption:
            self.text(s, caption, pic.left, pic.top + pic.height + Inches(0.05), pic.width, Inches(0.3), 11,
                      color=MUTED, align=PP_ALIGN.CENTER)

    def table(self, s, header, rows, x, y, w, col_w=None, size=13, highlight_row=None, row_h=0.42):
        shape = s.shapes.add_table(len(rows) + 1, len(header), x, y, w, Inches(row_h * (len(rows) + 1)))
        t = shape.table
        if col_w:
            for i, cw in enumerate(col_w):
                t.columns[i].width = Inches(cw)
        for j, h in enumerate(header):
            self.cell(t.cell(0, j), h, size, True, WHITE, ASPHALT)
        for i, row in enumerate(rows, 1):
            fill = RGBColor(0xFF, 0xF4, 0xCC) if highlight_row == i - 1 else (WHITE if i % 2 else LIGHT)
            for j, v in enumerate(row):
                self.cell(t.cell(i, j), str(v), size, j == 0, INK, fill)
        return t

    def cell(self, cell, text, size, bold, color, fill):
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = ""
        self.runs(p, text, size, bold, color, BODY)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    def bar_chart(self, s, categories, series, x, y, w, h, colors, title=None, maximum=100, legend=True, number_format='0'):
        data = CategoryChartData()
        data.categories = categories
        for name, values in series:
            data.add_series(name, values)
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, data)
        ch = gf.chart
        ch.has_title = bool(title)
        if title:
            ch.chart_title.text_frame.text = title
            self.style(ch.chart_title.text_frame.paragraphs[0].runs[0], 14, True, INK, BODY)
        ch.has_legend = legend
        if legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(12)
            ch.legend.font.name = BODY
        va = ch.value_axis
        va.maximum_scale = maximum
        va.minimum_scale = 0
        va.major_gridlines.format.line.color.rgb = LINE
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(11)
        va.tick_labels.font.color.rgb = MUTED
        ca = ch.category_axis
        ca.tick_labels.font.size = Pt(12)
        ca.tick_labels.font.color.rgb = INK
        ca.format.line.color.rgb = LINE
        plot = ch.plots[0]
        plot.gap_width = 60
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = number_format
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(10)
        dl.font.color.rgb = INK
        for sr, col in zip(plot.series, colors):
            sr.format.fill.solid()
            sr.format.fill.fore_color.rgb = col
        return ch


# --------------------------------------------------------------------------- #
def build(results_dir: Path, out: Path):
    R = Results(results_dir)
    D = Deck()
    m = R.models

    # 1 Title
    s = D.slide(dark=True, notes="Lakshya opens: introduce the team and the problem in one sentence. Each member presents their own model.")
    D.text(s, "ICT-4442 DEEP LEARNING · FINAL PRESENTATION", Inches(0.8), Inches(0.9), Inches(11), Inches(0.4), 14, True, YELLOW)
    D.text(s, "Road Damage Detection and Classification", Inches(0.8), Inches(1.5), Inches(11.5), Inches(0.9), 44, True, WHITE, HEAD)
    D.text(s, "Comparing four deep-learning families on RDD2022 for automated road-condition monitoring",
           Inches(0.8), Inches(2.45), Inches(11), Inches(0.6), 20, color=LINE)
    for i, (name, reg, key, owned) in enumerate(C.TEAM):
        x = Inches(0.8 + i * 3.0)
        D.card(s, x, Inches(4.2), Inches(2.8), Inches(1.75), ASPHALT2)
        D.badge(s, "ABCD"[i], x + Inches(0.2), Inches(4.4), MODEL_COLORS[key], 0.55)
        D.text(s, name.title() if name.isupper() else name, x + Inches(0.2), Inches(5.08), Inches(2.5), Inches(0.35), 15, True, WHITE)
        D.text(s, reg, x + Inches(0.2), Inches(5.4), Inches(2.5), Inches(0.3), 12, color=LINE)
        D.text(s, owned.split(" (")[0], x + Inches(0.95), Inches(4.45), Inches(1.75), Inches(0.5), 12, color=YELLOW,
               anchor=MSO_ANCHOR.MIDDLE)
    D.text(s, "School of Computer Engineering · MIT Manipal · " + C.MONTH_YEAR, Inches(0.8), Inches(6.5), Inches(10), Inches(0.4), 13, color=LINE)

    # 2 Problem
    s = D.slide(notes="Lakshya: why manual inspection fails; why automate; what RDD2022 gives us. Keep to ~1 min.")
    D.title(s, "Manual road surveys do not scale", "Problem & motivation")
    D.bullets(s, ["Slow and subjective|— visual inspection by crews, different inspectors disagree",
                  "Infrequent|— large networks are revisited only every few years",
                  "Costly when late|— a crack left untreated becomes a pothole: repair cost and accident risk grow",
                  "Opportunity|— dashcams and phones + deep learning allow continuous, low-cost monitoring"],
              Inches(0.6), Inches(1.8), Inches(6.4), Inches(4.5), 18, gap=14)
    for i, (big, small) in enumerate([("47,420", "road images"), ("6", "countries"), ("4", "damage classes"), (">55k", "labelled instances")]):
        x, y = Inches(7.5 + (i % 2) * 2.75), Inches(1.9 + (i // 2) * 2.2)
        D.card(s, x, y, Inches(2.5), Inches(1.9))
        D.text(s, big, x, y + Inches(0.3), Inches(2.5), Inches(0.9), 44, True, ASPHALT, HEAD, PP_ALIGN.CENTER)
        D.text(s, small, x, y + Inches(1.2), Inches(2.5), Inches(0.4), 14, color=MUTED, align=PP_ALIGN.CENTER)
    D.text(s, "RDD2022 (Arya et al., 2022) — Japan, India, Czech Republic, Norway, USA, China", Inches(7.5), Inches(6.35),
           Inches(5.3), Inches(0.3), 11, color=MUTED, italic=True)

    # 3 Objectives + classes
    counts = R.split_counts()
    s = D.slide(notes="Lakshya: three objectives; the four classes; the extra D44 class is dropped. Same split for every model.")
    D.title(s, "Objectives and the four damage classes", "Objectives & dataset")
    D.bullets(s, ["O1|Implement 4 models from distinct families (CNN, detector, segmenter, transformer)",
                  "O2|Benchmark all on ONE fixed RDD2022 test split with one common protocol",
                  "O3|Quantify speed and size → best model for edge / dashcam deployment"],
              Inches(0.6), Inches(1.75), Inches(12), Inches(1.6), 17, gap=8)
    classes = [("D00", "Longitudinal crack", "parallel to traffic"), ("D10", "Transverse crack", "across the lane"),
               ("D20", "Alligator crack", "fatigue mesh"), ("D40", "Pothole", "material loss")]
    for i, (code, name, desc) in enumerate(classes):
        x = Inches(0.6 + i * 3.08)
        D.card(s, x, Inches(3.75), Inches(2.85), Inches(2.6))
        D.text(s, code, x + Inches(0.25), Inches(3.95), Inches(2.4), Inches(0.6), 30, True, ASPHALT, HEAD)
        D.text(s, name, x + Inches(0.25), Inches(4.6), Inches(2.4), Inches(0.4), 16, True)
        D.text(s, desc, x + Inches(0.25), Inches(4.98), Inches(2.4), Inches(0.35), 13, color=MUTED)
        n = f"{sum(counts[2].get(sp, {}).get(code, 0) for sp in ('train', 'val', 'test')):,}" if counts else TBD
        D.text(s, f"{n} instances", x + Inches(0.25), Inches(5.6), Inches(2.4), Inches(0.4), 14, True, MODEL_COLORS["resnet50"])

    # 4 Pipeline
    s = D.slide(notes="Lakshya: explain the common protocol — the key idea that makes the comparison fair. Classifiers get the box; detector and segmenter must find it; a missed instance counts as wrong.")
    D.title(s, "One split, four families, one common protocol", "Method overview")
    D.picture(s, HERE / "assets" / "pipeline.png", Inches(0.6), Inches(1.6), Inches(12.1), Inches(4.4))
    D.text(s, "Every GT test instance must receive a class.  Classifier: crop → class.  YOLOv8: best-IoU box ≥ 0.5.  "
              "U-Net: majority class in box (≥ 5 % damage).  Missed = false negative.",
           Inches(0.6), Inches(6.2), Inches(12.1), Inches(0.6), 15, color=INK)

    # 5 Literature
    s = D.slide(notes="Kumar: what prior work shows — YOLO dominates real-time detection, U-Net with ResNet encoder for cracks, transformers more accurate but heavier; nobody compares families on the same split.")
    D.title(s, "What prior work tells us", "Literature")
    lit = [("Maeda et al. 2018", "SSD-MobileNet, 9k images", "first benchmark; rare classes poorly recalled"),
           ("Arya et al. 2021", "cross-country transfer", "Japan-only model: F1 ≤ 0.05 on India D00"),
           ("Wang et al. 2025 (YOLO-RD)", "YOLOv8 + attention modules", "mAP 25.75 % on RDD2022-Japan; small cracks hard"),
           ("Lau et al. 2020", "U-Net + ResNet-34 encoder", "F1 96 % CFD / 73 % Crack500"),
           ("Zhang & Zhang 2024", "TransUNet / SwinUNet vs CNN", "transformers best IoU (80.8 %), but slower"),
           ("Gap", "—", "no same-split comparison across families → our work")]
    D.table(s, ["Study", "Method", "Key finding"], [list(r) for r in lit], Inches(0.6), Inches(1.7), Inches(12.1),
            [3.3, 3.6, 5.2], 14, highlight_row=5, row_h=0.62)

    # 6-9 Model slides
    model_text = {
        "resnet50": ("ResNet-50 CNN classifier", "Model A · Lakshya Jain",
                     ["Architecture|50 layers, 16 bottleneck residual blocks; skip connections keep gradients alive",
                      "Input|GT box + 10 % context → 224×224 crop; ImageNet-pretrained, new 4-way head",
                      "Training|AdamW 2e-4, cosine, AMP, class-weighted CE + label smoothing 0.05, early stop",
                      "Augment|resized crop, h-flip, colour jitter — no rotation (would turn D00 into D10)",
                      "Why|standard CNN baseline: locality + translation equivariance suit crack texture"],
                     R.figure("resnet50", "confusion_matrix.png")),
        "yolov8": ("YOLOv8 one-stage detector", "Model B · Dhruv Maheshwari",
                   ["Architecture|CSPDarknet (C2f) backbone → PAN-FPN neck → anchor-free decoupled head",
                    "Loss|BCE (class) + CIoU + Distribution Focal Loss (box); NMS at inference",
                    "Training|YOLOv8s, COCO-pretrained, 640 px, mosaic, cosine LR, 30 % background images",
                    "Output|boxes + class on the full frame → also used for dashcam video (5 fps)",
                    "Why|operational framing: must find damage, real-time; dominates RDD challenges"],
                   R.figure("yolov8", "test", "confusion_matrix_normalized.png")),
        "unet": ("U-Net segmentation (ResNet-34 encoder)", "Model C · Kumar Kshitij",
                 ["Architecture|encoder–decoder with skip connections; ResNet-34 encoder (ImageNet)",
                  "Labels|RDD2022 has only boxes → box-derived masks (box pixels = class), small boxes on top",
                  "Training|512 px, weighted CE (bg 0.5) + Dice loss, AdamW + OneCycle, early stop on mIoU",
                  "Output|5-class pixel map → damage area (basis for severity)",
                  "Limitation|IoU measures damage regions, not exact crack outlines"],
                 R.figure("unet", "unet_samples.png")),
        "vit_b16": ("Vision Transformer ViT-B/16", "Model D · Manas Laxman Kudipady",
                    ["Architecture|224 px crop → 196 patches of 16×16 → 12 transformer layers, 12 heads",
                     "Idea|self-attention links every patch to every other → global context of long cracks",
                     "Training|same crops/split as ResNet-50; lr 5e-5, wd 0.05, 1-epoch warm-up + cosine",
                     "Trade-off|weak inductive bias → needs pretraining / data; 86 M parameters",
                     "Why|attention family; clean CNN-vs-transformer test on identical inputs"],
                    R.figure("vit_b16", "confusion_matrix.png")),
    }
    for idx, key in enumerate(MODEL_ORDER):
        name, kicker, items, fig = model_text[key]
        s = D.slide(notes=f"{PRESENTER[key]} presents: architecture, why this family, training choices, own results. "
                          f"Be ready for viva questions on loss, metrics and failure cases.")
        D.title(s, name, kicker)
        D.badge(s, "ABCD"[idx], Inches(12.05), Inches(0.5), MODEL_COLORS[key])
        D.bullets(s, items, Inches(0.6), Inches(1.75), Inches(6.6), Inches(4.4), 15, gap=10)
        c = R.common(key)
        mm = m.get(key, {})
        stats = [("Macro-F1", pct(c["macro_f1"]) + "%" if c else TBD), ("Accuracy", pct(c["accuracy"]) + "%" if c else TBD)]
        if key == "yolov8":
            d = mm.get("detection_test") or {}
            stats = [("mAP@0.5", num(d.get("mAP50")) if d else TBD), ("Macro-F1*", pct(c["macro_f1"]) + "%" if c else TBD)]
        if key == "unet":
            sg = mm.get("segmentation_test") or {}
            stats = [("mIoU (damage)", num(sg.get("mIoU_damage")) if sg else TBD), ("Macro-F1*", pct(c["macro_f1"]) + "%" if c else TBD)]
        stats.append(("Params / latency", f"{num(mm['params'] / 1e6, 1)} M · {num(mm.get('latency_ms_bs1'), 0)} ms" if mm else TBD))
        for j, (label, value) in enumerate(stats):
            x = Inches(0.6 + j * 2.25)
            D.card(s, x, Inches(5.75), Inches(2.1), Inches(1.05))
            D.text(s, value, x, Inches(5.82), Inches(2.1), Inches(0.55), 22 if j < 2 else 15, True, MODEL_COLORS[key], HEAD, PP_ALIGN.CENTER,
                   anchor=MSO_ANCHOR.MIDDLE)
            D.text(s, label, x, Inches(6.38), Inches(2.1), Inches(0.3), 11, color=MUTED, align=PP_ALIGN.CENTER)
        D.picture(s, fig, Inches(7.5), Inches(1.6), Inches(5.3), Inches(5.0))
        if key in ("yolov8", "unet"):
            D.text(s, "* common protocol: missed instances count as errors", Inches(7.5), Inches(6.65), Inches(5.3), Inches(0.3), 10,
                   color=MUTED, italic=True, align=PP_ALIGN.RIGHT)

    # 10 Phase-2 CNN evidence
    s = D.slide(notes="Lakshya: Phase-2 real runs — five CNN backbones on crops (validation split). Accuracy saturates ~77-79 %; MobileNetV3 is within 1.4 F1 points with 8x fewer params → crop classification is limited by labels/ambiguity, not capacity.")
    D.title(s, "Phase 2: CNN backbones saturate near 78 %", "Preliminary evidence (validation split)")
    cls = C.INTERIM["classifiers"]
    D.bar_chart(s, [c[0] for c in cls], [("Val accuracy (%)", [c[2] for c in cls]), ("Val macro-F1 (%)", [c[5] for c in cls])],
                Inches(0.6), Inches(1.6), Inches(8.2), Inches(5.0), [ASPHALT2, YELLOW], maximum=100, number_format='0.0')
    D.card(s, Inches(9.2), Inches(1.8), Inches(3.55), Inches(4.6))
    D.bullets(s, ["ConvNeXt-Tiny best|(74.3 % F1)", "MobileNetV3|within 1.4 pts of ResNet-50 at 8× fewer params",
                  "Takeaway|capacity is not the bottleneck — label ambiguity and box quality are"],
              Inches(9.45), Inches(2.05), Inches(3.1), Inches(4.2), 14, gap=12)

    # 11 Comparison table
    best = R.best()
    s = D.slide(notes="Manas: the single comparison table. Read the ranking, then explain the miss-rate column: classifiers are given the box, so their numbers measure recognition only.")
    D.title(s, "Common comparison on the same test instances", "Results")
    header = ["Model", "Params (M)", "Latency (ms)", "Accuracy %", "Macro-F1 %", "Miss %", "Native metric"]
    rows = [[R.row(k)[h] for h in ["Model", "Params (M)", "Latency (ms)", "Acc. (%)", "Macro-F1 (%)", "Miss (%)", "Native metric"]] for k in MODEL_ORDER]
    D.table(s, header, rows, Inches(0.6), Inches(1.75), Inches(12.1), [2.6, 1.3, 1.4, 1.4, 1.4, 1.1, 2.9], 15,
            highlight_row=MODEL_ORDER.index(best) if best and R.complete else None, row_h=0.62)
    n_inst = f"{R.final['n_common_instances']:,}" if R.final else TBD
    D.text(s, f"{n_inst} shared test instances · miss = GT damage not found by the localising model · native metric: mAP for YOLOv8, mIoU/Dice for U-Net",
           Inches(0.6), Inches(5.1), Inches(12.1), Inches(0.4), 13, color=MUTED)
    if R.complete:
        bc = R.common(best)
        msg = f"{R.name(best)} gives the best macro-F1 ({pct(bc['macro_f1'])} %)."
    else:
        msg = f"Best model: {TBD} (filled automatically after notebook 05)."
    D.card(s, Inches(0.6), Inches(5.65), Inches(12.1), Inches(0.95), RGBColor(0xFF, 0xF4, 0xCC))
    D.text(s, msg, Inches(0.9), Inches(5.65), Inches(11.5), Inches(0.95), 20, True, INK, HEAD, anchor=MSO_ANCHOR.MIDDLE)

    # 12 Per-class chart
    s = D.slide(notes="Manas: per-class F1. Point to the hardest class and to where detector/segmenter lose — thin cracks.")
    D.title(s, "Where each model wins and loses", "Per-class F1 (common protocol)")
    if all(R.common(k) for k in MODEL_ORDER):
        series = [(R.name(k), [100 * R.common(k)["per_class"][c]["f1"] for c in ("D00", "D10", "D20", "D40")]) for k in MODEL_ORDER]
        D.bar_chart(s, ["D00 longitudinal", "D10 transverse", "D20 alligator", "D40 pothole"], series,
                    Inches(0.6), Inches(1.6), Inches(8.4), Inches(5.1), [MODEL_COLORS[k] for k in MODEL_ORDER])
    else:
        D.picture(s, None, Inches(0.6), Inches(1.6), Inches(8.4), Inches(5.1))
    D.card(s, Inches(9.35), Inches(1.8), Inches(3.4), Inches(4.7))
    insight = []
    if R.complete:
        means = {c: sum(R.common(k)["per_class"][c]["f1"] for k in MODEL_ORDER) / 4 for c in ("D00", "D10", "D20", "D40")}
        insight.append(f"Hardest|{min(means, key=means.get)} (mean F1 {pct(min(means.values()))} %)")
        insight.append(f"Easiest|{max(means, key=means.get)} (mean F1 {pct(max(means.values()))} %)")
    else:
        insight += [f"Hardest class|{TBD}", f"Easiest class|{TBD}"]
    insight.append("Thin cracks|elongated boxes → IoU ≥ 0.5 is hard for the detector")
    D.bullets(s, insight, Inches(9.6), Inches(2.05), Inches(2.95), Inches(4.3), 14, gap=12)

    # 13 Why they differ + error analysis
    s = D.slide(notes="Dhruv + Kumar: four reasons for the differences, then show failure cases. This analysis is what is graded.")
    D.title(s, "Why performance differs", "Discussion & error analysis")
    reasons = [("1", "Task difficulty", "Classifiers get the box; YOLO/U-Net must find thin cracks in a full scene"),
               ("2", "Box geometry", "Elongated D00/D10 boxes: small shift → IoU < 0.5 → miss"),
               ("3", "Inductive bias", "CNN locality is data-efficient; ViT's global attention needs more data"),
               ("4", "Imbalance & domain", "D00 dominates; 6 countries differ in texture, camera and resolution")]
    for i, (k, head, body) in enumerate(reasons):
        y = Inches(1.7 + i * 1.22)
        D.badge(s, k, Inches(0.6), y + Inches(0.08), ASPHALT, 0.55)
        D.text(s, head, Inches(1.35), y, Inches(5.0), Inches(0.4), 17, True)
        D.text(s, body, Inches(1.35), y + Inches(0.4), Inches(5.1), Inches(0.7), 14, color=MUTED)
    gallery = R.figure("final", "failures_yolov8.png") or R.figure("final", "accuracy_by_box_size.png")
    D.picture(s, gallery, Inches(6.8), Inches(1.6), Inches(6.0), Inches(4.6),
              caption="Failure cases: diagonal cracks (D00↔D10), partial alligator patterns, shadowed potholes")

    # 14 Conclusion
    s = D.slide(dark=True, notes="Manas closes: answer the three objectives, then future work. Invite questions.")
    D.title(s, "Conclusion & future work", "Takeaways", dark=True)
    if R.complete:
        y = m["yolov8"]
        concl = [f"Most accurate|{R.name(best)} — macro-F1 {pct(R.common(best)['macro_f1'])} % on the common protocol",
                 f"Most deployable|YOLOv8: end-to-end on full frames, {num(y.get('latency_ms_bs1'), 0)} ms/image, {num(y['params'] / 1e6, 1)} M params",
                 "Bottleneck|localising thin cracks, not recognising the damage type"]
    else:
        concl = [f"Most accurate|{TBD}", "Most deployable|YOLOv8 — the only end-to-end, real-time model on full frames",
                 "Bottleneck|localising thin cracks, not recognising the damage type"]
    D.bullets(s, concl, Inches(0.6), Inches(1.8), Inches(6.3), Inches(4), 17, WHITE, 14)
    D.card(s, Inches(7.3), Inches(1.8), Inches(5.45), Inches(4.4), ASPHALT2)
    D.text(s, "Next", Inches(7.6), Inches(2.0), Inches(5), Inches(0.4), 18, True, YELLOW, HEAD)
    D.bullets(s, ["Real crack masks (Crack500 / SAM-refined) for U-Net", "Detector + classifier two-stage pipeline",
                  "Transformer encoders (TransUNet, Swin)", "INT8 export + Jetson field test on dashcam video"],
              Inches(7.6), Inches(2.6), Inches(4.9), Inches(3.4), 15, WHITE, 10)
    D.text(s, "Thank you — questions?", Inches(0.6), Inches(6.2), Inches(6), Inches(0.5), 22, True, YELLOW, HEAD)

    D.prs.save(str(out))
    print("Saved", out, "| slides:", D.n, "| complete results:", R.complete)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--results", type=Path, default=HERE.parent / "results")
    parser.add_argument("--out", type=Path, default=HERE.parent / "ICT4442_Final_Presentation.pptx")
    args = parser.parse_args()
    build(args.results, args.out)
